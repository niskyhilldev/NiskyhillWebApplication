from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import pandas as pd
import re

# Load the Excel file
file_path = "The Data.xlsx"
# wb = load_workbook(file_path)
# ws = wb.active  # You can also specify a sheet by name: wb["Sheet1"]

# # Define a fill style (yellow in this case)
# highlight = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

# Highlight a specific cell, e.g., B2
# ws["B2"].fill = highlight

df = pd.read_excel(file_path)

# cnt = 0
# for value in df["First"]:
#     if pd.notnull(value):  # Skip empty cells
#         if re.search(r'[^A-Za-z]', str(value)):  # Check for non-alphabetic characters
#             cnt += 1
#             print(value)
            #child, stillborn, infant, son, baby

# print(f"Total non-alphabetic entries in the first column: {cnt}")

def remove_titles(column_name):
    titles = ["Dr\\.", "Rev\\.", "Mrs\\.", "Mr\\.", "Jr\\."]
    pattern = re.compile(r"^(" + "|".join(titles) + r")\s*", re.IGNORECASE)


    # Iterate over the column and print matches
    for index, value in df[column_name].items():
        if pd.notnull(value) and pattern.match(str(value)):
            print(value)
            new_value = pattern.sub("", str(value))
            df.at[index, column_name] = new_value

def separate_middle_names(column_name):
    exclude_words = ["stillborn", "baby", "son", "daughter", "child", "infant"]

    # Iterate and print matching values
    for index, value in df[column_name].items():
        if pd.notnull(value):
            value_str = str(value).strip()  # Remove leading/trailing whitespace
            if value_str == "Maxilond {Maxine)":
                continue
            # Check first character is alphabetical
            if value_str and value_str[0].isalpha():
                value_lower = value_str.lower()
                if " " in value_str and not any(word in value_lower for word in exclude_words):
                    print(f"Row {index + 2}: {value_str}")
                    parts = value_str.split(" ", 1)
                    first_name = parts[0]
                    middle_name = parts[1] if len(parts) > 1 else ""

                    # Update dataframe
                    df.at[index, "First"] = first_name
                    df.at[index, "Middle"] = middle_name
                        
                # new_value = pattern.sub("", str(value))
                # df.at[index, column_name] = new_value

def remove_leading_special_characters(column_name):
    # Iterate and print matching values

    for index, value in df[column_name].items():
        if pd.notnull(value):
            value_str = str(value).strip()  # Remove leading/trailing whitespace
            if value_str == "(Frank) Bruce":
                continue
            if re.match(r'[^A-Za-z]', value_str):  # starts with a non-alphabetical character
                # Remove everything up to the first alphabetical character
                new_value = re.sub(r'^[^A-Za-z]*', '', value_str)
                print(f"Row {index + 1}: {new_value}")
                # Optionally, update the dataframe
                df.at[index, column_name] = new_value

def differs_by_one_char(s1, s2):
    """Return True if s1 and s2 differ by exactly one character."""
    if s1 is None or s2 is None:
        return False
    s1, s2 = str(s1), str(s2)
    # If lengths differ by more than 1, can't differ by exactly one char
    if abs(len(s1) - len(s2)) > 1:
        return False
    # Check for exact one-character difference
    differences = 0
    i, j = 0, 0
    while i < len(s1) and j < len(s2):
        if s1[i] != s2[j]:
            differences += 1
            if differences > 1:
                return False
            # If lengths are equal, move both pointers
            if len(s1) == len(s2):
                i += 1
                j += 1
            # If s1 is longer, skip a character in s1
            elif len(s1) > len(s2):
                i += 1
            else:  # s2 is longer
                j += 1
        else:
            i += 1
            j += 1
    # Account for an extra char at the end
    if i < len(s1) or j < len(s2):
        differences += 1
    return differences == 1

def typos(column_name):
    previous_value = None
    for index, value in df[column_name].items():
        if pd.notnull(value) and previous_value is not None:
            if differs_by_one_char(previous_value, value):
                print(f"Row {index + 1}: {value} (differs from previous: {previous_value})")
        previous_value = value


column_name = "Last"
for index, value in df[column_name].items():
    if pd.notnull(value):
        value_str = str(value).strip()
        # Check if it contains any non-alphabetic character
        if re.search(r'[^A-Za-z]', value_str):
            print(f"Row {index + 1}: {value_str}")

# while True:
#     try:
#         # Save the modified DataFrame to a new Excel file
#         df.to_excel("The Data.xlsx", index=False)
#         break
#     except PermissionError:
#         input("Please close 'The Data.xlsx' if it is open and press Enter to retry...")
